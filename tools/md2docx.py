"""把 Markdown 转成 Word (.docx)。用法：python -m tools.md2docx <in.md> <out.docx>"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def _add_runs(p, text: str) -> None:
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            p.add_run(part)


def convert(md_path: str, docx_path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for lvl in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            doc.styles[lvl].font.name = "微软雅黑"
            doc.styles[lvl].element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        except Exception:  # noqa: BLE001
            pass

    for line in Path(md_path).read_text(encoding="utf-8").splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:].strip(), level=1)
        elif s.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, s[2:])
        elif re.match(r"^\s*[-*] ", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, re.sub(r"^\s*[-*] ", "", s))
        elif re.match(r"^\s*\d+\. ", s):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\s*\d+\. ", "", s))
        elif s.strip() == "---":
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            _add_runs(p, s)
    doc.save(docx_path)


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    convert(src, dst)
    print("OK ->", dst)
