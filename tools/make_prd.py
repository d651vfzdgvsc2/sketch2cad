"""生成 Sketch2CAD 的 PRD（Word 文档）到桌面。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
OUT = DESKTOP / "Sketch2CAD：图像到CAD图纸的多Agent矢量化系统.docx"


def main() -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def h(text, level=1):
        doc.add_heading(text, level=level)

    def p(text):
        doc.add_paragraph(text)

    def pb(text):
        para = doc.add_paragraph()
        para.add_run(text).bold = True
        return para

    def b(text):
        doc.add_paragraph(text, style="List Bullet")

    def num(text):
        doc.add_paragraph(text, style="List Number")

    title = doc.add_heading("", level=0)
    r = title.add_run("Sketch2CAD：图像到CAD图纸的多Agent矢量化系统")
    r.font.name = "微软雅黑"
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    sub = doc.add_paragraph("产品需求文档（PRD）")
    sub.runs[0].bold = True
    sub.alignment = 1

    doc.add_paragraph("版本 v1.0　|　作者：潘昕　|　文档类型：产品需求文档（PRD）")

    h("一、项目概述", 1)
    p("Sketch2CAD 是一款把「黑白线稿 / 图片」自动转换为「可编辑 DXF CAD 图纸」的辅助级矢量化系统。"
      "它面向工程图与假山（有机形状）两类场景，采用「算法清洗 + 精确矢量化 + 多 Agent 协同 + "
      "渲染回图客观仲裁」的技术路线。")
    p("定位：不做“一键全自动出完美 CAD”，而是由 AI 产出 70~80% 的初稿，再由人工精修，"
      "从而把“把图片重画成 CAD”这件重复劳动降到最低。")

    h("二、背景与问题", 1)
    b("人工描图：把图片转成 CAD 是完全重复劳动，慢、贵、易错。")
    b("传统光栅转矢量工具：只会描像素边界，输出成千上万条碎线段，不懂“这是直线 / 圆 / 标注”。")
    b("纯大模型：能“看懂图”，但给不出精确坐标，画出来“大概像、尺寸对不上”。")
    b("业界现状：尚无“一键全自动出完美 CAD”的方案，因此本项目做“辅助级”落地。")

    h("三、目标与非目标", 1)
    pb("目标：")
    b("实现图片 → 可编辑 DXF 的自动化输出，覆盖工程图与假山双场景。")
    b("用客观指标（SSIM / IoU 等）量化质量，并能持续优化。")
    b("结果可控、可解释、可回退（“保最优”机制，输出永不劣于基线）。")
    pb("非目标：")
    b("不追求全自动完美重建（定位为辅助级）。")
    b("不做 3D 建模 / 3D→2D 投影。")
    b("不面向照片级彩色图纸，聚焦黑白线稿。")

    h("四、目标用户", 1)
    b("制图 / 工程技术人员：把纸质图、截图、扫描件数字化为可编辑 CAD。")
    b("园林景观设计师：把假山、置石线稿转为 CAD 轮廓。")
    b("需要把图片转 CAD 的办公 / 设计用户。")

    h("五、范围", 1)
    b("✅ 支持：简单 / 中等工程图（方块、圆、槽、孔、圆弧、椭圆、少量标注）；假山等有机曲线。")
    b("⚠️ 受限：复杂多视图装配图（图框 + 标题栏 + 剖面线 + 透视角）目前仅能“部分重建”。")

    h("六、功能需求（FR）", 1)
    p("FR1　模式选择：用户手动选择「工程图」或「假山」，避免自动误判。")
    p("FR2　图像清洗：二值化去噪、图框 / 标题栏剔除、多方向剖面线剔除、尺寸线剔除。")
    p("FR3　精确矢量化：直线、圆（圆周覆盖度验证，抑制假圆）、圆弧（Kasa 拟合）、椭圆（拟合 + "
      "覆盖度验证）；并做图元装配（共线合并 / 角度吸附 / 端点吸附 / 矩形装配）。")
    p("FR4　多 Agent 协同：识别 Agent（VLM + OCR）、提取 Agent、校验 Agent、决策 Agent，"
      "用 LangGraph 编排。")
    p("FR5　多臂提案：A（确定性 CV）、MV（多视图分割）、D（接地式代码生成）、"
      "B（纯语义代码生成）、DS（DeepSeek 直接看图写脚本）。")
    p("FR6　客观仲裁：把每个提案渲染回图、与原图比对（SSIM / IoU / Chamfer），取最优，"
      "并保留历史最优（“保最优”）。")
    p("FR7　比例尺校准：用 OCR 到的尺寸标注反推 mm/px，导出“真实尺寸 DXF”。")
    p("FR8　假山通道：用 Potrace 对有机轮廓做平滑曲线拟合，输出多段线/样条。")
    p("FR9　输出：可编辑 DXF 文件 + 预览图 + 指标报告。")
    p("FR10　图形界面：选择图片、选择模式、一键处理、进度日志、自动打开结果。")

    h("七、非功能需求（NFR）", 1)
    b("性能：假山模式约 2 秒（纯本地）；工程图模式约 1~2 分钟（含云端 API）。")
    b("成本：使用通义千问 VLM / DeepSeek，token 用量可控。")
    b("可靠性：脚本子进程沙箱执行 + 错误反馈重试 + 保最优回退。")
    b("可用性：图形界面一键操作，无需命令行。")
    b("可扩展：引擎与场景解耦，新增图元类型 / 模式不影响主干。")

    h("八、系统架构", 1)
    p("整体为“清洗 → 矢量化 → 多臂提案 → 客观仲裁 → 输出”的流水线；工程图走多 Agent 集成，"
      "假山走 Potrace 通道。（详细流程见配套流程图《Sketch2CAD流程图》。）")

    h("九、技术栈", 1)
    b("语言 / 图像：Python、OpenCV、scikit-image、NumPy、SciPy")
    b("矢量化：HoughLinesP、HoughCircles（+覆盖度验证）、Kasa 圆/弧拟合、椭圆拟合、图元装配")
    b("CAD：ezdxf（生成 DXF）、matplotlib（DXF 渲染回图）")
    b("OCR：RapidOCR（本地、支持旋转文字）")
    b("大模型：通义千问 qwen-vl-max（视觉理解）、DeepSeek（看图写代码 / 决策）")
    b("编排：LangGraph；界面：tkinter")

    h("十、评测指标与结果", 1)
    b("图像级：SSIM、IoU、Dice、Chamfer；结构级：图元数量 / 类型。")
    b("简单 / 中等工程图：SSIM 0.75~0.84；复杂装配图：较基线 0.38 → 0.63；"
      "多臂仲裁平均最佳 0.756。")
    b("比例尺校准：real01 标注 60×50，还原包围盒 61.5×50.0 mm，误差 <3%。")

    h("十一、里程碑", 1)
    num("阶段 0：合成数据生成器 + 评测脚本（IoU/SSIM/Chamfer）。")
    num("阶段 1：OpenCV 确定性矢量化基线跑通。")
    num("阶段 2：接入 OCR + VLM 语义理解。")
    num("阶段 3：LangGraph 多 Agent 协同 + 校验回环。")
    num("阶段 4：批量对比评测 + 报告。")
    num("阶段 5：算法增强（图框/剖面线/尺寸线剔除、圆验证、椭圆、图元装配）。")
    num("阶段 6：A/B/C/D/E + DS 多臂对照实验 + 集成仲裁固化。")
    num("阶段 7：双模式（工程图 / 假山）+ GUI + 比例尺校准。")

    h("十二、风险与限制", 1)
    b("复杂多视图图的瓶颈是“检测不全”，图元装配只能规整“已抓到的”。")
    b("依赖云端 API：有成本与延迟，结果有随机性（由“保最优”兜底）。")
    b("样条 / 自由曲线尚未支持（椭圆已支持）。")

    h("十三、成果", 1)
    p("相关成果已申请国家发明专利；实现图片到可编辑 DXF 的自动化输出，"
      "简单 / 中等图 SSIM 0.75~0.84，复杂图较基线 0.38→0.63。")

    doc.save(str(OUT))
    print("OK ->", OUT)


if __name__ == "__main__":
    main()
